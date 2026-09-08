# -*- coding: utf-8 -*-
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from rows_data import ROWS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT_NAME = "Arial"

wb = openpyxl.Workbook()

# ---------- helper styles ----------
header_fill = PatternFill("solid", fgColor="1F3864")
header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
title_font = Font(name=FONT_NAME, size=16, bold=True, color="1F3864")
subtitle_font = Font(name=FONT_NAME, size=11, italic=True, color="595959")
wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
wrap_center = Alignment(wrap_text=True, vertical="center", horizontal="center")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

DECISION_COLORS = {
    "KEEP": "FFF2CC",       # yellow-ish -> must keep
    "SKIP": "D9EAD3",       # green -> automatable
    "INSPECT": "CFE2F3",    # blue -> reduced physical
    "ACT": "F4CCCC",        # red-ish -> demand based / action trigger
}
def decision_color(text):
    t = (text or "").upper()
    for key in ("KEEP", "SKIP", "INSPECT", "ACT"):
        if t.startswith(key):
            return DECISION_COLORS[key]
    return "FFFFFF"

def freq_category(equip, task, source):
    pass

# assign strategic classification per row using notes already encoded via automatable/decision text
def classify(r):
    if r["decision"].startswith("KEEP"):
        return "C — רגולציה/ביטוח/בטיחות (לא ניתן לביטול)"
    if "מבוסס תפוסה" in r["decision"] or "מבוסס ביקוש" in r["decision"] or "Demand" in r["notes"]:
        return "D — מבוסס תפוסה/ביקוש (Demand-Based)"
    if "לא נדרש" in r["sensor"]:
        return "A — נתון קיים בבקר (אוטומציה מיידית)"
    return "B — נדרש חיישן IoT נוסף"

# =========================================================
# SHEET 1: מפת החלטות מאסטר
# =========================================================
ws = wb.active
ws.title = "מפת החלטות מאסטר"
ws.sheet_view.rightToLeft = True

headers = [
    "#", "מערכת", "ציוד / רכיב", "משימה (מקור: תוכנית האחזקה)", "תדירות נוכחית",
    "סיווג אסטרטגי", "מופעים לשנה (היום)", "מה נבדק בפועל (תמצית)", "מקור החובה (Compliance)",
    "נתון קיים בבקר/BMS?", "חיישן IoT נדרש", "בדיקה פיזית חובה עפ\"י דין/תקן?",
    "ניתן למעבר לניהול מבוסס חריגות?", "החלטת ברירת מחדל", "תנאי SKIP", "תנאי INSPECT", "תנאי ACT",
    "זמן ביצוע היום (דק')", "שעות/שנה - היום", "מופעים חזויים לשנה",
    "זמן ביצוע חזוי (דק')", "שעות/שנה - חזוי", "חיסכון שעות/שנה", "% חיסכון", "הערות / הנחות",
]

title = "Master Decision Matrix — מפת החלטות תפעוליות לבניין (KEEP / SKIP / INSPECT / ACT)"
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
c = ws.cell(row=1, column=1, value=title)
c.font = title_font
c.alignment = wrap
ws.row_dimensions[1].height = 26

subtitle = ("מקור: תוכנית האחזקה המונעת (230 משימות, 12 מערכות) — נבחרו 36 המשימות בעלות פוטנציאל ההשפעה הגבוה ביותר לשלב 1. "
            "כל שעות/זמן הן הערכות ראשוניות (MVP) לאימות מול נתוני אמת בפיילוט.")
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
c = ws.cell(row=2, column=1, value=subtitle)
c.font = subtitle_font
c.alignment = wrap
ws.row_dimensions[2].height = 30

header_row = 3
for j, h in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=j, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_center
    cell.border = border
ws.row_dimensions[header_row].height = 46

start_data_row = header_row + 1
for i, r in enumerate(ROWS):
    row = start_data_row + i
    strategic = classify(r)
    values = [
        i + 1,
        r["system"],
        r["equip"],
        r["task"],
        r["freq"],
        strategic,
        r["occ_now"],
        r["what"],
        r["source"],
        r["bms"],
        r["sensor"],
        r["phys_mandatory"],
        r["automatable"],
        r["decision"],
        r["skip"],
        r["inspect"],
        r["act"],
        r["time_now"],
        None,  # formula
        r["occ_future"],
        r["time_future"],
        None,  # formula
        None,  # formula
        None,  # formula
        r["notes"],
    ]
    for j, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=j, value=v)
        cell.font = Font(name=FONT_NAME, size=10)
        cell.alignment = wrap
        cell.border = border

    # formulas
    ws.cell(row=row, column=19, value=f"=G{row}*R{row}/60")   # hours now = occ_now * time_now /60
    ws.cell(row=row, column=22, value=f"=T{row}*U{row}/60")   # hours future = occ_future*time_future/60
    ws.cell(row=row, column=23, value=f"=S{row}-V{row}")      # savings hours
    ws.cell(row=row, column=24, value=f"=IFERROR(W{row}/S{row},0)")  # % savings

    # number formats
    ws.cell(row=row, column=7).number_format = "#,##0"
    ws.cell(row=row, column=18).number_format = "#,##0"
    ws.cell(row=row, column=19).number_format = "#,##0.0"
    ws.cell(row=row, column=20).number_format = "#,##0"
    ws.cell(row=row, column=21).number_format = "#,##0"
    ws.cell(row=row, column=22).number_format = "#,##0.0"
    ws.cell(row=row, column=23).number_format = "#,##0.0"
    ws.cell(row=row, column=24).number_format = "0%"

    # decision color
    fill = PatternFill("solid", fgColor=decision_color(r["decision"]))
    ws.cell(row=row, column=14).fill = fill
    ws.cell(row=row, column=6).fill = PatternFill("solid", fgColor="F2F2F2")

    ws.row_dimensions[row].height = 60

end_data_row = start_data_row + len(ROWS) - 1
total_row = end_data_row + 1
ws.cell(row=total_row, column=1, value="סה\"כ").font = Font(name=FONT_NAME, bold=True)
ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
for col, letter in [(19, "S"), (22, "V"), (23, "W")]:
    cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{start_data_row}:{letter}{end_data_row})")
    cell.font = Font(name=FONT_NAME, bold=True)
    cell.number_format = "#,##0.0"
    cell.fill = PatternFill("solid", fgColor="D9D9D9")
avg_pct = ws.cell(row=total_row, column=24, value=f"=IFERROR(W{total_row}/S{total_row},0)")
avg_pct.font = Font(name=FONT_NAME, bold=True)
avg_pct.number_format = "0%"
avg_pct.fill = PatternFill("solid", fgColor="D9D9D9")
for col in range(1, len(headers) + 1):
    ws.cell(row=total_row, column=col).border = border

# column widths
widths = {
    1: 5, 2: 14, 3: 20, 4: 26, 5: 12, 6: 28, 7: 12, 8: 34, 9: 30, 10: 22, 11: 22,
    12: 20, 13: 20, 14: 22, 15: 30, 16: 30, 17: 30, 18: 12, 19: 12, 20: 14, 21: 12,
    22: 12, 23: 12, 24: 10, 25: 34,
}
for col, w in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

ws.freeze_panes = f"A{start_data_row}"
ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{end_data_row}"

# =========================================================
# SHEET 2: תקציר מנהלים (Dashboard)
# =========================================================
ws2 = wb.create_sheet("תקציר מנהלים")
ws2.sheet_view.rightToLeft = True
ws2.merge_cells("A1:E1")
c = ws2.cell(row=1, column=1, value="BUILDING OPERATIONS — תקציר מנהלים (שלב 1: Blueprint)")
c.font = title_font
ws2.row_dimensions[1].height = 26

ws2.merge_cells("A2:E2")
c = ws2.cell(row=2, column=1, value="מבוסס על 36 המשימות הראשונות שנותחו מתוך 230 המשימות בתוכנית האחזקה. נתוני MVP להערכה — לאימות בפיילוט.")
c.font = subtitle_font

sheet_name = "'מפת החלטות מאסטר'"
keep_count_formula = f"=SUMPRODUCT(--(LEFT({sheet_name}!N{start_data_row}:N{end_data_row},4)=\"KEEP\"))"
kpis = [
    ("סה\"כ משימות שנותחו בשלב 1", len(ROWS), "#,##0", None),
    ("מתוכן ניתנות למעבר מלא/חלקי לניהול מבוסס-חריגות (SKIP/INSPECT/ACT)",
     f"={len(ROWS)}-({keep_count_formula[1:]})", "#,##0", None),
    ("מתוכן חייבות להישאר KEEP (רגולציה / ביטוח / בטיחות חיים)",
     keep_count_formula, "#,##0", None),
    ("שעות עבודה בשנה — היום (36 המשימות)", f"={sheet_name}!S{total_row}", "#,##0", "שעות"),
    ("שעות עבודה בשנה — לאחר מעבר לניהול מבוסס-חריגות", f"={sheet_name}!V{total_row}", "#,##0", "שעות"),
    ("חיסכון שעות עבודה לשנה", f"={sheet_name}!W{total_row}", "#,##0", "שעות"),
    ("אחוז חיסכון ממוצע (36 המשימות שנותחו)", f"={sheet_name}!X{total_row}", "0%", None),
]

row = 4
ws2.cell(row=row, column=1, value="מדד").font = Font(name=FONT_NAME, bold=True)
ws2.cell(row=row, column=2, value="ערך").font = Font(name=FONT_NAME, bold=True)
for col in (1, 2):
    ws2.cell(row=row, column=col).fill = header_fill
    ws2.cell(row=row, column=col).font = header_font
row += 1
for label, formula, fmt, unit in kpis:
    ws2.cell(row=row, column=1, value=label).alignment = wrap
    ws2.cell(row=row, column=1).font = Font(name=FONT_NAME, size=11)
    v = ws2.cell(row=row, column=2, value=formula)
    v.number_format = fmt
    v.font = Font(name=FONT_NAME, size=12, bold=True, color="1F3864")
    v.alignment = Alignment(horizontal="center")
    if unit:
        ws2.cell(row=row, column=3, value=unit).font = Font(name=FONT_NAME, size=10, italic=True)
    ws2.row_dimensions[row].height = 22
    row += 1

row += 1
ws2.cell(row=row, column=1, value="פילוח לפי סיווג אסטרטגי").font = Font(name=FONT_NAME, bold=True, size=12)
row += 1
ws2.cell(row=row, column=1, value="סיווג").font = Font(name=FONT_NAME, bold=True)
ws2.cell(row=row, column=2, value="מס' משימות").font = Font(name=FONT_NAME, bold=True)
ws2.cell(row=row, column=3, value="חיסכון שעות/שנה").font = Font(name=FONT_NAME, bold=True)
for col in (1, 2, 3):
    ws2.cell(row=row, column=col).fill = header_fill
    ws2.cell(row=row, column=col).font = header_font
row += 1
classes = [
    "A — נתון קיים בבקר (אוטומציה מיידית)",
    "B — נדרש חיישן IoT נוסף",
    "C — רגולציה/ביטוח/בטיחות (לא ניתן לביטול)",
    "D — מבוסס תפוסה/ביקוש (Demand-Based)",
]
for cls in classes:
    ws2.cell(row=row, column=1, value=cls).alignment = wrap
    ws2.cell(row=row, column=2,
             value=f"=COUNTIF({sheet_name}!F{start_data_row}:F{end_data_row},\"{cls}\")").alignment = Alignment(horizontal="center")
    ws2.cell(row=row, column=3,
             value=f"=SUMIF({sheet_name}!F{start_data_row}:F{end_data_row},\"{cls}\",{sheet_name}!W{start_data_row}:W{end_data_row})").number_format = "#,##0.0"
    ws2.row_dimensions[row].height = 30
    row += 1

ws2.column_dimensions["A"].width = 55
ws2.column_dimensions["B"].width = 16
ws2.column_dimensions["C"].width = 18

# =========================================================
# SHEET 3: מקרא ומתודולוגיה
# =========================================================
ws3 = wb.create_sheet("מקרא ומתודולוגיה")
ws3.sheet_view.rightToLeft = True
ws3.column_dimensions["A"].width = 32
ws3.column_dimensions["B"].width = 95

def add_section(row, title_text):
    ws3.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c = ws3.cell(row=row, column=1, value=title_text)
    c.font = Font(name=FONT_NAME, bold=True, size=13, color="1F3864")
    return row + 1

def add_pair(row, k, v):
    ws3.cell(row=row, column=1, value=k).font = Font(name=FONT_NAME, bold=True, size=10)
    ws3.cell(row=row, column=1).alignment = wrap
    c = ws3.cell(row=row, column=2, value=v)
    c.font = Font(name=FONT_NAME, size=10)
    c.alignment = wrap
    ws3.row_dimensions[row].height = max(18, 15 * (len(v) // 90 + 1))
    return row + 1

r = 1
ws3.merge_cells("A1:B1")
ws3.cell(row=1, column=1, value="מקרא ומתודולוגיה — Master Decision Matrix").font = title_font
r = 3

r = add_section(r, "1. מטרת המסמך")
r = add_pair(r, "",
    "שלב 1 (Blueprint) בתהליך בניית מנוע ההחלטות התפעוליות: פירוק תוכנית האחזקה המונעת הקיימת (230 משימות, 12 מערכות) "
    "וניסוח לוגיקת KEEP / SKIP / INSPECT / ACT לכל משימה, כבסיס לאפיון תוכנה ולשיחה עם מפתח.")
r += 1

r = add_section(r, "2. ארבע ההחלטות האפשריות")
r = add_pair(r, "KEEP", "המשימה נשארת כפי שהיא, בתדירותה המקורית — כי מקור החובה שלה (תקן/רגולציה/ביטוח/בטיחות חיים) אינו מאפשר החלפה בטכנולוגיה, גם אם הטכנולוגיה קיימת.")
r = add_pair(r, "SKIP", "המשימה מבוטלת כברירת מחדל ומוחלפת בניטור דיגיטלי רציף — מבוצעת פיזית רק כאשר מתקבלת חריגה.")
r = add_pair(r, "INSPECT", "תדירות הביקור הפיזי מופחתת (לא מבוטלת), לרוב משילוב של ניטור דיגיטלי חלקי + הוספת חיישן IoT.")
r = add_pair(r, "ACT", "המשימה עוברת ממודל 'לוח זמנים קבוע' למודל 'מבוסס ביקוש/תפוסה/מילוי' — מבוצעת כשנדרש בפועל, לא לפי יום קבוע מראש.")
r += 1

r = add_section(r, "3. הסיווג האסטרטגי (A/B/C/D)")
r = add_pair(r, "A — נתון קיים בבקר", "המידע הדרוש להחלטה כבר זמין היום מהבקר/ה-BMS הקיים. פוטנציאל האוטומציה הגבוה והמהיר ביותר — לא דורש השקעת חומרה.")
r = add_pair(r, "B — נדרש חיישן IoT נוסף", "הנתון אינו קיים כיום (למשל רעד, דליפה, הפרש לחץ). נדרשת השקעה נקודתית בחיישן זול לפני שניתן לצמצם ביקורים פיזיים.")
r = add_pair(r, "C — רגולציה/ביטוח/בטיחות", "מקור החובה חיצוני ומחייב (תקן, חוק, דרישת מבטח, בטיחות חיים) — המנוע אינו מציע SKIP גם אם קיימת טכנולוגיה, זהו 'שער' ה-Compliance.")
r = add_pair(r, "D — מבוסס תפוסה/ביקוש", "המשימה קשורה לשימוש בפועל בשטח (ניקיון, פינוי אשפה, סיורים) — עוברת ממודל תדיר קבוע למודל המגיב לנתוני תפוסה/מילוי בפועל.")
r += 1

r = add_section(r, "4. הסבר עמודות מפת ההחלטות")
cols_explain = [
    ("מערכת / ציוד / משימה", "כפי שמופיעים בתוכנית האחזקה המקורית (12 לשוניות, 230 משימות) — שומר על שקיפות מלאה מול המקור."),
    ("תדירות נוכחית / מופעים לשנה (היום)", "התדירות כפי שנקבעה היום בתוכנית האחזקה, מתורגמת למספר מופעים שנתי."),
    ("מה נבדק בפועל", "תמצית הצ'ק-ליסט המקורי לאותה משימה."),
    ("מקור החובה (Compliance)", "הסיבה שבגללה המשימה קיימת: תקן ישראלי, חוק, הוראות יצרן, דרישת ביטוח, SLA או נוהל פנימי. זהו הבסיס לכל החלטת KEEP."),
    ("נתון קיים בבקר/BMS?", "האם המידע הדרוש להערכת מצב המערכת כבר מגיע היום מבקר/רכזת/תוכנת ניהול קיימת."),
    ("חיישן IoT נדרש", "אם המידע חסר — איזה חיישן זול נדרש כדי להשלים אותו (רטט, דליפה, מילוי, הפרש לחץ וכו')."),
    ("בדיקה פיזית חובה עפ\"י דין/תקן?", "דגל Compliance מפורש — אם כן, לא מוצעת החלטת SKIP למשימה, ללא תלות ביכולת הטכנולוגית."),
    ("ניתן למעבר לניהול מבוסס חריגות?", "הערכה כללית (כן/חלקי/לא) של מידת ההתאמה של המשימה למודל Condition-Based."),
    ("החלטת ברירת מחדל", "אחת מארבע ההחלטות (KEEP/SKIP/INSPECT/ACT), עם צביעה תואמת."),
    ("תנאי SKIP / INSPECT / ACT", "הלוגיקה בפועל: מתי המנוע מדלג, מתי הוא שולח לבדיקה מופחתת, ומתי הוא פותח קריאת שירות/פעולה."),
    ("זמן ביצוע (דק') / מופעים חזויים / שעות בשנה", "הערכות MVP לחישוב חיסכון פוטנציאלי — ראו הנחות עבודה בסעיף 5."),
]
for k, v in cols_explain:
    r = add_pair(r, k, v)
r += 1

r = add_section(r, "5. הנחות עבודה ומגבלות (MVP)")
assumptions = [
    "אומדני 'זמן ביצוע' ו'מופעים חזויים' הם הערכות סבירות של צוות הפיתוח בשלב Blueprint, ולא נתוני אמת נמדדים — מטרתן לבנות מודל חישוב שניתן להזין אליו בהמשך נתוני זמן אמיתיים (למשל ממגדלי הארבעה).",
    "כל שורת KEEP מייצגת בכוונה חיסכון אפס או נמוך — זו דוגמה מכוונת לכך שהמנוע אינו ממליץ לבטל משימה רק כי קיימת טכנולוגיה, כאשר מקור החובה חיצוני ומחייב.",
    "שורות 'סיור ביטחון' מסומנות באזהרה מפורשת: כל צמצום מותנה באישור מראש של גורם הביטחון/הביטוח של הבניין — לא רק ביכולת הטכנולוגית, בהתאם לעיקרון הזהירות שהוגדר במסמך החזון.",
    "המשימות שנבחרו (36 מתוך 230) הן אלה בעלות פוטנציאל ההשפעה הגבוה ביותר (תדירות גבוהה + פוטנציאל אוטומציה) — הרחבה ליתר המשימות היא שלב 2 של התהליך.",
    "מקור החובה (Compliance) מבוסס על ידע כללי בתחום התקינה בישראל (ת\"י 1220 ודומיו) ואינו מהווה ייעוץ משפטי/הנדסי — יש לאמת מול יועץ בטיחות/חשמל/כיבוי אש מוסמך לפני יישום בפועל.",
]
for a in assumptions:
    r = add_pair(r, "•", a)
r += 1

r = add_section(r, "6. השלבים הבאים (מתוך מסמך החזון)")
next_steps = [
    "הרחבת המטריצה ליתר 194 המשימות בתוכנית האחזקה (ולתוכנית הניקיון, כשתתקבל).",
    "אימות מול נתוני אמת ממגדלי הארבעה — זמני ביצוע בפועל, מפת חיישנים/בקרים קיימת בפועל בכל מגדל.",
    "בניית Decision Engine V1 היברידי: שכבת Rules (Compliance) + שכבת Analytics/ML (זיהוי חריגות) + שכבת Generative AI (הסברים והמלצות).",
    "בניית מסך Dashboard יומי ('BUILDING OPERATIONS — TODAY') המציג KEEP/SKIP/INSPECT/ACT לכל בניין בזמן אמת.",
]
for s in next_steps:
    r = add_pair(r, "→", s)

ws.sheet_state = "visible"
ws2.sheet_state = "visible"
ws3.sheet_state = "visible"
wb.active = 0

out_path = Path(__file__).resolve().parent.parent / "output" / "master_decision_matrix.xlsx"
out_path.parent.mkdir(parents=True, exist_ok=True)
wb.save(out_path)
print("saved", out_path)
